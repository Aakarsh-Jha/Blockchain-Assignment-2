import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def create_submission_document(
    student_name="Aakarsh Jha",
    reg_num="230905248",
    assignment_title="Blockchain Technology (CSE4444) - Assignment 2",
    dapp_title="CertiTrust: Decentralized Academic Credential Verification System",
    github_link="https://github.com/Aakarsh-Jha/Blockchain-Assignment-2",
    deployed_address="0x742d35Cc6634C0532925a3b844Bc454e4438f44e",
    network="Ganache Local / Remix VM / Sepolia Testnet"
):
    # ---------------- 1. Generate DOCX ----------------
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("CSE4444 - Blockchain Technology\nAssignment 2 Submission Report")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph()

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    data = [
        ("Student Name:", student_name),
        ("Registration Number:", reg_num),
        ("Assignment Title:", assignment_title),
        ("Title of Developed DApp:", dapp_title),
        ("GitHub Repository Link:", github_link),
    ]

    col_widths = [Inches(2.4), Inches(4.5)]
    for i, (label, val) in enumerate(data):
        row = table.rows[i]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width, cell_val.width = col_widths[0], col_widths[1]
        
        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(11)

        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(11)
        if label.startswith("GitHub"):
            r_val.font.color.rgb = RGBColor(37, 99, 235)
            r_val.font.underline = True

        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{"F8FAFC" if i % 2 == 0 else "FFFFFF"}"/>')
        cell_lbl._tc.get_or_add_tcPr().append(shading)
        shading2 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{"F8FAFC" if i % 2 == 0 else "FFFFFF"}"/>')
        cell_val._tc.get_or_add_tcPr().append(shading2)

    doc.add_paragraph()

    h1 = doc.add_heading(level=1)
    h1.add_run("1. DApp Overview & Problem Statement").font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "CertiTrust is an original decentralized application (DApp) designed to address academic credential fraud. "
        "Academic certificates and degrees are cryptographically registered by an authorized institution on the Ethereum blockchain. "
        "Each certificate record contains the student's name, degree title, document reference hash, and an immutable block timestamp. "
        "Employers, verification agencies, and students can instantly query the blockchain to verify authenticity without relying on slow manual registrar confirmations."
    )

    h2 = doc.add_heading(level=1)
    h2.add_run("2. Technical Stack & Architecture").font.color.rgb = RGBColor(30, 41, 59)

    tech_items = [
        ("Smart Contract Language:", "Solidity (^0.8.20)"),
        ("Development & Compilation:", "Remix IDE"),
        ("Target Blockchain Network:", network),
        ("Contract Address:", deployed_address),
        ("Frontend Technologies:", "HTML5, CSS3 (Modern Glassmorphic UI), JavaScript"),
        ("Web3 Integration:", "Ethers.js v6"),
        ("Wallet Integration:", "MetaMask Extension (Account & Chain Event Listeners)")
    ]

    tech_table = doc.add_table(rows=len(tech_items), cols=2)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, (k, v) in enumerate(tech_items):
        r = tech_table.rows[idx]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width, c1.width = Inches(2.5), Inches(4.4)
        c0.paragraphs[0].add_run(k).font.bold = True
        c1.paragraphs[0].add_run(v)

    doc.add_paragraph()

    h3 = doc.add_heading(level=1)
    h3.add_run("3. Input Validation & Test Scenario Verification").font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph(
        "The smart contract implements comprehensive access control and validation using Solidity modifiers and require statements. "
        "The application was tested across three critical test vectors:"
    )

    test_matrix = [
        ("TC-01 (Valid Input)", "Owner Account", "Valid student name, degree, and hash", "Transaction confirmed; new certificate ID emitted and stored on-chain."),
        ("TC-02 (Invalid Input)", "Owner Account", "Empty student name string (\"\")", "Transaction reverts with 'Validation Error: Student name cannot be empty.'"),
        ("TC-03 (Unauthorized Input)", "Non-Owner (Account 2)", "Attempt to execute issueCertificate()", "Transaction reverts with 'Error: Unauthorized! Only the university authority can perform this action.'")
    ]

    t_table = doc.add_table(rows=len(test_matrix) + 1, cols=4)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Test Case", "Caller", "Test Input", "Observed EVM Behavior"]
    for j, h in enumerate(headers):
        cell = t_table.rows[0].cells[j]
        run = cell.paragraphs[0].add_run(h)
        run.font.bold = True
        run.font.size = Pt(10)
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E2E8F0"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    for i, (tc, caller, inp, res) in enumerate(test_matrix):
        row = t_table.rows[i + 1]
        for col_idx, text in enumerate([tc, caller, inp, res]):
            cell = row.cells[col_idx]
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(9.5)

    doc.add_paragraph()

    # Section 4: Evidence Screenshots
    h4 = doc.add_heading(level=1)
    h4.add_run("4. Screenshots & Execution Evidence").font.color.rgb = RGBColor(30, 41, 59)

    screenshots_info = [
        ("Figure 1: Smart Contract Compilation (Remix IDE)", "screenshots/01_contract_compilation.png"),
        ("Figure 2: Smart Contract Deployment (Remix VM / Network)", "screenshots/02_contract_deployment.png"),
        ("Figure 3: MetaMask Wallet Connection Handshake", "screenshots/03_metamask_connection.png"),
        ("Figure 4: Successful Transaction Execution & Mining", "screenshots/04_successful_transaction.png"),
        ("Figure 5: Failed Transaction Revert (Validation / Unauthorized)", "screenshots/05_failed_transaction.png"),
        ("Figure 6: Blockchain Credential Retrieval & Verification in UI", "screenshots/06_blockchain_data_display.png"),
    ]

    for caption, img_path in screenshots_info:
        if os.path.exists(img_path):
            p_cap = doc.add_paragraph()
            r_cap = p_cap.add_run(caption)
            r_cap.font.bold = True
            r_cap.font.size = Pt(10)
            doc.add_picture(img_path, width=Inches(6.0))
            doc.add_paragraph() # Spacing

    safe_name = "".join(c for c in student_name if c.isalnum() or c in (" ", "_")).strip().replace(" ", "_")
    safe_reg = "".join(c for c in reg_num if c.isalnum() or c in ("-", "_")).strip()
    docx_filename = f"{safe_name}_{safe_reg}_BT_Assignment2.docx"
    pdf_filename = f"{safe_name}_{safe_reg}_BT_Assignment2.pdf"

    doc.save(docx_filename)
    print(f"Generated Word Document: {docx_filename}")

    # ---------------- 2. Generate PDF via ReportLab ----------------
    doc_pdf = SimpleDocTemplate(
        pdf_filename,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    story = []
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        alignment=1,
        textColor=colors.HexColor('#0F172A')
    )
    story.append(Paragraph("<b>Blockchain Technology (CSE4444)</b><br/>Assignment 2 Submission Document", title_style))
    story.append(Spacer(1, 12))

    t_data = [
        ["Student Name:", student_name],
        ["Registration Number:", reg_num],
        ["Assignment Title:", assignment_title],
        ["Title of Developed DApp:", dapp_title],
        ["GitHub Repository Link:", github_link]
    ]
    rl_table = Table(t_data, colWidths=[150, 390])
    rl_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#0F172A')),
        ('FONTNAME', (0,0), (0,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(rl_table)
    story.append(Spacer(1, 14))

    story.append(Paragraph("<b>1. DApp Overview & Problem Statement</b>", styles['Heading2']))
    story.append(Paragraph(
        "CertiTrust is an original decentralized application (DApp) designed to address academic credential fraud. "
        "Academic certificates and degrees are cryptographically registered by an authorized institution on the Ethereum blockchain. "
        "Each certificate record contains the student's name, degree title, document reference hash, and an immutable block timestamp. "
        "Employers, verification agencies, and students can instantly query the smart contract to verify authenticity.",
        styles['Normal']
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>2. Input Validation & Test Cases Summary</b>", styles['Heading2']))
    tc_data = [
        ["Test Case", "Caller", "Inputs", "Observed Outcome"],
        ["TC-01 (Valid)", "Owner", "Valid Name & Course", "Mined & issued successfully (ID #1)."],
        ["TC-02 (Invalid)", "Owner", "Empty student name \"\"", "Reverted: 'Student name cannot be empty.'"],
        ["TC-03 (Unauthorized)", "Account 2", "issueCertificate() call", "Reverted: 'Unauthorized! Only authority can perform this action.'"]
    ]
    t_tc = Table(tc_data, colWidths=[90, 75, 140, 235])
    t_tc.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 8.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
    ]))
    story.append(t_tc)
    story.append(Spacer(1, 14))

    story.append(Paragraph("<b>3. Screenshots & Evidence</b>", styles['Heading2']))
    for caption, img_path in screenshots_info:
        if os.path.exists(img_path):
            img_story = [
                Paragraph(f"<b>{caption}</b>", styles['Normal']),
                Spacer(1, 4),
                Image(img_path, width=480, height=230),
                Spacer(1, 12)
            ]
            story.append(KeepTogether(img_story))

    doc_pdf.build(story)
    print(f"Generated PDF Document: {pdf_filename}")

if __name__ == "__main__":
    name = sys.argv[1] if len(sys.argv) > 1 else "Aakarsh Jha"
    reg = sys.argv[2] if len(sys.argv) > 2 else "230905248"
    repo = sys.argv[3] if len(sys.argv) > 3 else "https://github.com/Aakarsh-Jha/Blockchain-Assignment-2"
    create_submission_document(student_name=name, reg_num=reg, github_link=repo)
